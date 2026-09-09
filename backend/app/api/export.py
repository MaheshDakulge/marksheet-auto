from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from app.core.database import get_supabase
from app.core.security import get_current_teacher
from supabase import Client
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import io

router = APIRouter(prefix="/export", tags=["export"])

# Colours
HEADER_BG = "1F4E79"
HEADER_FG = "FFFFFF"
SUBHDR_BG = "2E75B6"
ALT_BG    = "EBF3FB"
PASS_BG   = "E2EFDA"
FAIL_BG   = "FFE0E0"
ATKT_BG   = "FFF2CC"


def _border():
    s = Side(style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)


def _hcell(ws, row, col, value, bg=HEADER_BG, fg=HEADER_FG, size=10):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=True, color=fg, size=size)
    c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = _border()
    return c


def _dcell(ws, row, col, value, bold=False, align="center", bg=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=10)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    c.border = _border()
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    return c


def _status_bg(status):
    s = (status or "").lower()
    if s == "pass": return PASS_BG
    if s == "fail": return FAIL_BG
    if s == "atkt": return ATKT_BG
    return None


@router.get("/{project_id}")
async def export_project(
    project_id: str,
    format: str = Query(default="excel", enum=["csv", "excel"]),
    teacher: dict = Depends(get_current_teacher),
    supabase: Client = Depends(get_supabase),
):
    """
    Export all students in this project as Excel or CSV.
    - 1 sheet only: Student Results
    - One row per student
    - Dynamic subject columns based on uploaded marksheets
    - Works correctly whether teacher uploaded 1 PDF or 200 PDFs
    - Re-exports always show the latest data
    - Analytics are in the Flutter frontend, NOT in this file
    """
    proj = (
        supabase.table("projects").select("*")
        .eq("id", project_id).eq("teacher_id", teacher["id"]).execute()
    )
    if not proj.data:
        raise HTTPException(status_code=404, detail="Project not found")
    project = proj.data[0]

    results = (
        supabase.table("student_results")
        .select("*, subject_marks(*), students(name, seat_no, prn)")
        .eq("project_id", project_id)
        .order("created_at")       # Sr No = order student was added
        .execute()
    )
    if not results.data:
        raise HTTPException(status_code=404, detail="No results found for this project")

    data = results.data

    # Collect all unique subjects across all students in insertion order
    subject_cols = []    # list of (course_code, course_name)
    seen_codes   = set()
    for r in data:
        for m in sorted(r.get("subject_marks", []), key=lambda x: x.get("sr_no") or 0):
            code = m.get("course_code") or f"Sub{m.get('sr_no', '')}"
            if code not in seen_codes:
                subject_cols.append((code, m.get("course_name", code)))
                seen_codes.add(code)

    # ── CSV ────────────────────────────────────────────────────────────────────
    if format == "csv":
        rows = []
        for sr_no, r in enumerate(data, 1):
            student = r.get("students") or {}
            row = {
                "Sr No":         sr_no,
                "PRN":           r.get("student_prn", ""),
                "Seat No":       student.get("seat_no", ""),
                "Name":          student.get("name", ""),
                "Examination":   r.get("examination", ""),
                "Semester":      r.get("semester", ""),
                "Total Credits": r.get("total_credits", ""),
                "Total EGP":     r.get("total_egp", ""),
                "SGPA":          r.get("sgpa", ""),
                "CGPA":          r.get("cgpa", ""),
                "Status":        (r.get("status") or "").upper(),
            }
            marks_by_code = {
                (m.get("course_code") or f"Sub{m.get('sr_no', '')}"): m
                for m in r.get("subject_marks", [])
            }
            for code, _ in subject_cols:
                m = marks_by_code.get(code, {})
                row[f"{code} Grade"] = m.get("grade_obtained", "")
                row[f"{code} GP"]    = m.get("grade_point", "")
                row[f"{code} EGP"]   = m.get("earned_gp", "")
            rows.append(row)

        buf = io.StringIO()
        pd.DataFrame(rows).to_csv(buf, index=False)
        buf.seek(0)
        fname = project.get("name", "export").replace(" ", "_")
        return StreamingResponse(
            iter([buf.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{fname}.csv"'},
        )

    # ── Excel ──────────────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Results"
    ws.sheet_view.showGridLines = False

    total_cols = 11 + len(subject_cols) * 3

    # Row 1 — Project title bar
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(total_cols, 11))
    title = (
        f"{project.get('name', '').upper()}  |  "
        f"{project.get('examination', '')}  |  "
        f"Semester: {project.get('semester', '')}"
    )
    t = ws.cell(row=1, column=1, value=title)
    t.font = Font(bold=True, size=13, color=HEADER_FG)
    t.fill = PatternFill("solid", fgColor=HEADER_BG)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    # Row 2 — Program + student count
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(total_cols, 11))
    info = f"Program: {project.get('program', '')}  |  Total Students: {len(data)}"
    s = ws.cell(row=2, column=1, value=info)
    s.font = Font(italic=True, size=10, color=HEADER_FG)
    s.fill = PatternFill("solid", fgColor=SUBHDR_BG)
    s.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 16

    # Row 3 — Fixed column headers + subject merged headers
    fixed_headers = [
        "Sr No", "PRN", "Seat No", "Name",
        "Examination", "Semester",
        "Total Credits", "Total EGP", "SGPA", "CGPA", "Status",
    ]
    for col, h in enumerate(fixed_headers, 1):
        _hcell(ws, 3, col, h)

    col = 12
    for code, name in subject_cols:
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 2)
        _hcell(ws, 3, col, f"{code}\n{name}", bg=SUBHDR_BG, size=9)
        col += 3
    ws.row_dimensions[3].height = 32

    # Row 4 — Subject sub-headers: Grade | GP | EGP
    for i in range(1, 12):
        _hcell(ws, 4, i, "", bg=SUBHDR_BG)
    col = 12
    for _ in subject_cols:
        _hcell(ws, 4, col,     "Grade", bg=SUBHDR_BG, size=9)
        _hcell(ws, 4, col + 1, "GP",    bg=SUBHDR_BG, size=9)
        _hcell(ws, 4, col + 2, "EGP",   bg=SUBHDR_BG, size=9)
        col += 3
    ws.row_dimensions[4].height = 16

    # Data rows — one row per student, starting at row 5
    for sr_no, r in enumerate(data, 1):
        row     = 4 + sr_no
        student = r.get("students") or {}
        status  = r.get("status", "")
        rbg     = ALT_BG if sr_no % 2 == 0 else None   # alternate row shading

        _dcell(ws, row, 1,  sr_no,                         bold=True,  bg=rbg)
        _dcell(ws, row, 2,  r.get("student_prn", ""),      align="left", bg=rbg)
        _dcell(ws, row, 3,  student.get("seat_no", ""),    bg=rbg)
        _dcell(ws, row, 4,  student.get("name", ""),       bold=True, align="left", bg=rbg)
        _dcell(ws, row, 5,  r.get("examination", ""),      align="left", bg=rbg)
        _dcell(ws, row, 6,  r.get("semester", ""),         bg=rbg)
        _dcell(ws, row, 7,  r.get("total_credits", ""),    bg=rbg)
        _dcell(ws, row, 8,  r.get("total_egp", ""),        bg=rbg)
        _dcell(ws, row, 9,  r.get("sgpa", ""),             bold=True, bg=rbg)
        _dcell(ws, row, 10, r.get("cgpa", ""),             bg=rbg)
        _dcell(ws, row, 11, status.upper(),                bold=True, bg=_status_bg(status))

        marks_by_code = {
            (m.get("course_code") or f"Sub{m.get('sr_no', '')}"): m
            for m in r.get("subject_marks", [])
        }
        col = 12
        for code, _ in subject_cols:
            m = marks_by_code.get(code, {})
            _dcell(ws, row, col,     m.get("grade_obtained", ""), bold=True, bg=rbg)
            _dcell(ws, row, col + 1, m.get("grade_point", ""),    bg=rbg)
            _dcell(ws, row, col + 2, m.get("earned_gp", ""),      bg=rbg)
            col += 3
        ws.row_dimensions[row].height = 16

    # Column widths
    fixed_widths = [6, 16, 10, 28, 28, 12, 13, 11, 8, 8, 10]
    for i, w in enumerate(fixed_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    col = 12
    for _ in subject_cols:
        ws.column_dimensions[get_column_letter(col)].width     = 8
        ws.column_dimensions[get_column_letter(col + 1)].width = 6
        ws.column_dimensions[get_column_letter(col + 2)].width = 8
        col += 3

    # Freeze top 4 rows so headers stay visible when scrolling 200 students
    ws.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = project.get("name", "export").replace(" ", "_")
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'},
    )